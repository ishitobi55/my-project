"""旧「売上管理表」→「仕入れ管理表 Ver.2.0」移行。

  count モード: 棚卸しリスト（実在庫を記入してもらう表）とタイトル変更リストを作る
    python3 build_v2.py count <旧表のCSVフォルダ> <出力xlsx>
  final モード: 記入済みの棚卸しリストから、新しい表に貼る移行データを作る
    python3 build_v2.py final <旧表のCSVフォルダ> <記入済み棚卸しリストxlsx> <出力xlsx> <発注日 YYYY/MM/DD>
"""
import csv
import re
import sys
from collections import Counter, OrderedDict

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

RECENT_TABS = ['売上2026.7', '売上2026.8', '売上2026.9', '売上2026.10']
YELLOW = PatternFill('solid', fgColor='FFF7CC')
GRAY = PatternFill('solid', fgColor='E5E5E5')
HEAD = PatternFill('solid', fgColor='D9E8F5')
BOLD = Font(bold=True)


def num(s):
    try:
        return float(str(s).replace(',', '').replace('¥', '').strip())
    except ValueError:
        return None


def norm(code):
    return code.strip().upper()


def load(src):
    def read(path):
        with open(path, encoding='utf-8') as fh:
            return list(csv.reader(fh))

    products = OrderedDict()
    for r in read(f'{src}/tabs/673925804.csv')[1:]:          # 商品リスト
        code = norm(r[0])
        if not code or code == '例':
            continue
        products[code] = {'code': code, 'name': r[1].strip() or f'（商品名未登録）{code}',
                          'price': num(r[2]), 'cost': num(r[4]), 'ship': num(r[5])}
    inv = {}
    for r in read(f'{src}/tabs/994680164.csv')[2:]:          # 在庫管理
        code = norm(r[3])
        if code and code != '例':
            inv[code] = num(r[6]) or 0
    sold_all, sold_recent, last_cost = Counter(), Counter(), {}
    tabs = [line.rstrip('\n').split(' ', 1) for line in open(f'{src}/tabs.txt', encoding='utf-8')]
    # 古い月 → 新しい月の順に読んで、最後に見た仕入値を「直近の仕入値」にする
    def month_key(t):
        m = re.search(r'(?:2026\.)?(\d+)月?\s*$', t[1].strip())
        return (1 if '2026' in t[1] else 0, int(m.group(1)) if m else 0)

    for gid, tab in sorted(tabs, key=month_key):
        if '売上' not in tab or '見本' in tab:
            continue
        for r in read(f'{src}/tabs/{gid}.csv')[1:]:
            if len(r) > 11 and r[2].strip() and r[3].strip():
                code = norm(r[2])
                sold_all[code] += 1
                if tab in RECENT_TABS:
                    sold_recent[code] += 1
                if num(r[11]) is not None:
                    last_cost[code] = num(r[11])
    return products, inv, sold_all, sold_recent, last_cost


def write_sheet(ws, header, data, widths, fills=None):
    ws.append(header)
    for c in ws[1]:
        c.font, c.fill = BOLD, HEAD
        c.alignment = Alignment(wrap_text=True, vertical='center')
    for d in data:
        ws.append(d)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for col, fill in (fills or {}).items():
        for row in ws.iter_rows(min_row=2, min_col=col, max_col=col):
            row[0].fill = fill
    ws.freeze_panes = 'A2'


def title_with_code(name, code):
    parts = re.split(r'([ 　])', name, maxsplit=1)
    if len(parts) == 3:
        return f'{parts[0]} {code} {parts[2]}'
    return f'{name} {code}'


def count_mode(src, out):
    products, inv, sold_all, sold_recent, last_cost = load(src)
    groups = OrderedDict()
    for p in products.values():
        groups.setdefault(p['name'], []).append(p)
    rows, titles = [], []
    for name, items in groups.items():
        active = [p for p in items if inv.get(p['code'], 0) > 0 or sold_recent[p['code']] > 0]
        if not active:
            continue
        head = max(active, key=lambda p: (sold_recent[p['code']], inv.get(p['code'], 0)))
        others = [p['code'] for p in items if p is not head]
        rows.append([head['code'], name, ', '.join(others),
                     sum(max(inv.get(p['code'], 0), 0) for p in items),
                     sum(sold_recent[p['code']] for p in items),
                     head['cost'], round(last_cost[head['code']], 1) if head['code'] in last_cost else None, head['price'], head['ship'], None])
        titles.append([head['code'], name, title_with_code(name, head['code'])])
    rows.sort(key=lambda r: r[0])
    titles.sort(key=lambda r: r[0])

    wb = Workbook()
    ws = wb.active
    ws.title = '説明'
    for line in [
        '■ 棚卸しリスト（実在庫の記入をお願いします）',
        '',
        '今の売上管理表の「在庫管理」タブの在庫数は、2026年の売上が反映されていませんでした。',
        '（例：A-29 は2026年に342個売れていますが、在庫管理では「販売7個・在庫9個」のまま）',
        'そのため、実際の在庫数を数えて「棚卸し」シートの黄色い列（J列）に入れてください。',
        '',
        '・在庫が0の商品は 0 を入れてください（空欄のままでもOKです。0として扱います）',
        '・リストに無い商品で在庫があるものは、いちばん下の行に商品番号・商品名・実在庫を追加してください',
        '・同じ商品名で番号が複数あるもの（C列）は、まとめて1行にしています。全部合わせた数を入れてください',
        '・D列（在庫管理の在庫数）は参考です。あまり当てになりません',
        '・E列は 2026年7月〜今日 に売れた個数です。数え漏れを防ぐ目安にしてください',
        '・F列は「商品リスト」の仕入れ値、G列は売上表で最後に使われた仕入値です。',
        '　違う場合は、新しい表にはF列（商品リスト）の値を入れる予定です。G列を使いたいものがあれば教えてください',
        '',
        '記入が終わったら、このファイルをチャットに送ってください。新しい表に貼る移行データを作ります。',
    ]:
        ws.append([line])
    ws['A1'].font = Font(bold=True, size=13)
    ws.column_dimensions['A'].width = 110

    write_sheet(wb.create_sheet('棚卸し'),
                ['商品番号', '商品名', 'まとめた番号', '在庫管理の\n在庫数(参考)', '7月以降の\n販売数',
                 '仕入れ値\n(商品リスト)', '直近の仕入値\n(売上表)', '販売価格', '発送送料', '★実在庫\n（ここに記入）'],
                rows, [9, 58, 14, 11, 10, 11, 11, 9, 9, 12], fills={10: YELLOW})
    write_sheet(wb.create_sheet('タイトル変更例'),
                ['商品番号', '今の商品名', 'タイトル例（最初の言葉の後ろに番号）'],
                titles, [9, 58, 66])
    wb.save(out)
    print('count rows', len(rows))


def final_mode(src, filled, out, order_date):
    products, inv, sold_all, sold_recent, last_cost = load(src)
    ws = load_workbook(filled, data_only=True)['棚卸し']
    rows, merged = [], []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r[0]:
            continue
        code, name, others, stock = norm(str(r[0])), r[1], r[2] or '', num(r[9] if r[9] is not None else 0) or 0
        if stock <= 0:
            continue
        codes = [code] + [norm(c) for c in str(others).split(',') if c.strip()]
        # 加重平均：各番号の（全期間の販売数＋在庫）を仕入数とみなして重みにする
        w_sum = q_sum = 0.0
        for c in codes:
            p = products.get(c)
            if not p or p['cost'] is None:
                continue
            q = sold_all[c] + max(inv.get(c, 0), 0)
            w_sum += q * p['cost']
            q_sum += q
        p = products.get(code, {})
        unit = round(w_sum / q_sum, 2) if q_sum else (p.get('cost') or num(r[5]))
        rep = ''
        if len(codes) > 1:
            rep = f'R{len({m[0] for m in merged}) + 1}'
            for c in codes:
                q = sold_all[c] + max(inv.get(c, 0), 0)
                merged.append([rep, c, name, q, products.get(c, {}).get('cost'), unit])
        rows.append([code, name or p.get('name'), rep, int(stock), p.get('price') or num(r[7]), unit,
                     p.get('ship') or num(r[8])])
    total = round(sum(r[3] * r[5] for r in rows), 2)
    ymd = order_date.replace('/', '')

    wb = Workbook()
    ws = wb.active
    ws.title = '説明'
    for line in [
        '■ 移行データの貼り方',
        '',
        '【1】「仕入管理（貼付用）」の2行目 A2〜I2 をコピー → 新しい表「仕入管理」の2行目 A2 に「値のみ貼り付け」（Ctrl+Shift+V）',
        '【2】「商品情報（貼付用）」の B〜C列 → 新しい表「商品情報」の B2 に「値のみ貼り付け」',
        '【3】「商品情報（貼付用）」の F〜R列 → 新しい表「商品情報」の F2 に「値のみ貼り付け」',
        '　　※ D列（商品管理番号）・E列（商品画像）は自動計算なので貼らない（グレー）',
        '【4】新しい表の「検算」がすべてOK、「在庫一覧」の数が実物と合うかを確認',
        '',
        f'・発注日：{order_date}（商品管理番号は「商品番号-{ymd}」）／為替レート 1／送料・手数料・関税 0',
        '・商品代(元/個) ＝ 1個あたりの原価（円）。為替レート1なので、そのまま原価になります',
        '・同じ商品名で番号が複数あったものは加重平均（重み＝全期間の販売数＋在庫数）',
        '・「設定」のPayPal掛け率は 1 のままにしてください',
        '',
        f'■ 商品 {len(rows)} 行 ／ 在庫 {sum(r[3] for r in rows)} 個 ／ 在庫原価合計 {total:,.2f} 円',
    ]:
        ws.append([line])
    ws['A1'].font = Font(bold=True, size=13)
    ws.column_dimensions['A'].width = 110

    write_sheet(wb.create_sheet('仕入管理（貼付用）'),
                ['発注日', '注文日', '配送番号', '為替レート\n(円/元)', '商品代合計\n(元)', '中国国内送料合計\n(元)',
                 'ラクマート手数料\n(円)', '国際配送料\n(円)', '関税\n(円)'],
                [[order_date, order_date, '移行在庫', 1, total, 0, 0, 0, 0]],
                [12, 12, 12, 11, 14, 15, 14, 12, 9], fills={i: YELLOW for i in range(1, 10)})
    write_sheet(wb.create_sheet('商品情報（貼付用）'),
                ['行番号', '商品番号', '発注日', '商品管理番号', '商品画像', '商品名', '色・サイズ', '訳あり記載',
                 'リピート番号', '販売開始日', '注文数', '販売価格\n(円)', '1688URL', '画像リンク', '出品URL',
                 '商品代\n(元/個)', '中国国内送料\n(元)', '発送送料\n(円)'],
                [['', r[0], order_date, '', '', r[1], '', '', r[2], order_date, r[3], r[4], '', '', '', r[5], 0, r[6]]
                 for r in rows],
                [6, 9, 11, 14, 8, 58, 9, 8, 8, 11, 7, 9, 9, 9, 9, 10, 10, 9],
                fills={1: GRAY, 4: GRAY, 5: GRAY})
    write_sheet(wb.create_sheet('リピート統合の内訳'),
                ['リピート番号', '商品番号', '商品名', '重み(販売数+在庫)', '仕入れ値', '加重平均単価'],
                merged, [10, 9, 58, 14, 10, 12])
    wb.save(out)
    print('final rows', len(rows), 'units', sum(r[3] for r in rows), 'total', total, 'merged', len(merged))


if __name__ == '__main__':
    if sys.argv[1] == 'count':
        count_mode(sys.argv[2], sys.argv[3])
    else:
        final_mode(sys.argv[2], sys.argv[3], sys.argv[4], sys.argv[5])
